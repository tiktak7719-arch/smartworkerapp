import zipfile
from pathlib import Path

from docx import Document


p = Path(r"D:\My Web Sites\smartworkerapp\Document_from_Hamad_Khan_Revised_Clickable_TOC.docx")
d = Document(p)
print("paragraphs", len(d.paragraphs))
print("inline_shapes", len(d.inline_shapes))
print("tables", len(d.tables))
print("sections", len(d.sections))
print("toc", sum(x.text.strip() == "Table of Contents" for x in d.paragraphs))
print("appendix", sum("APPENDIX A" in x.text for x in d.paragraphs))
with zipfile.ZipFile(p) as z:
    xml = z.read("word/document.xml").decode("utf-8", "ignore")
    print("toc_links", xml.count('w:hyperlink w:anchor="toc_'))
    print("media", len([x for x in z.namelist() if x.startswith("word/media/")]))
    print("zip", z.testzip())
