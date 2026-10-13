from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\Laptop Valley\Downloads\Document from Hamad Khan.docx")
FINAL = Path(r"D:\My Web Sites\smartworkerapp\Document_from_Hamad_Khan_Revised_Clickable_TOC.docx")
DOWNLOADS = Path(r"C:\Users\Laptop Valley\Downloads")


def q(tag: str) -> str:
    return qn(tag)


def normalise(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().casefold()


def set_tnr(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Times New Roman"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(q(key), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(document: Document) -> None:
    sizes = {
        "Normal": 12,
        "Body Text": 12,
        "List Paragraph": 12,
        "Normal (Web)": 12,
        "Heading 1": 16,
        "Heading 2": 14,
        "Heading 3": 16,
        "Heading 4": 13,
        "Heading 5": 12,
        "Title": 12,
        "Header": 9,
        "Footer": 9,
    }
    for name, size in sizes.items():
        if name not in [style.name for style in document.styles]:
            continue
        style = document.styles[name]
        style.font.name = "Times New Roman"
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(q(key), "Times New Roman")
        style.font.size = Pt(size)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)


def format_paragraph(paragraph, *, toc_entry=False, table_header=False) -> None:
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    style = paragraph.style.name
    if toc_entry:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0)
        paragraph.paragraph_format.first_line_indent = Inches(0)
    elif table_header:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif style.startswith("Heading") or style == "Title":
        # Keep title/heading layout intact; it should not be stretched by justification.
        if paragraph.alignment is None:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif paragraph.alignment not in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in paragraph.runs:
        set_tnr(run)


def format_everything(document: Document) -> None:
    configure_styles(document)
    for paragraph in document.paragraphs:
        format_paragraph(paragraph)
    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    format_paragraph(paragraph, table_header=(row_index == 0))
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_tnr(run, 9)


def add_bookmark(paragraph, name: str, identifier: int) -> None:
    element = paragraph._p
    start = OxmlElement("w:bookmarkStart")
    start.set(q("w:id"), str(identifier))
    start.set(q("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(q("w:id"), str(identifier))
    element.insert(1 if element.pPr is not None else 0, start)
    element.append(end)


def make_hyperlink(paragraph, target: str, text: str, indent_inches: float) -> None:
    element = paragraph._p
    for child in list(element):
        if child.tag != q("w:pPr"):
            element.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(indent_inches)
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.line_spacing = 1.5
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(q("w:anchor"), target)
    hyperlink.set(q("w:history"), "1")
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(q(key), "Times New Roman")
    run_properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(q("w:val"), "24")
    run_properties.append(size)
    colour = OxmlElement("w:color")
    colour.set(q("w:val"), "000000")
    run_properties.append(colour)
    underline = OxmlElement("w:u")
    underline.set(q("w:val"), "none")
    run_properties.append(underline)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    element.append(hyperlink)


def add_page_break(paragraph) -> None:
    run = paragraph.add_run()
    br = OxmlElement("w:br")
    br.set(q("w:type"), "page")
    run._r.append(br)


def heading_level(paragraph) -> int | None:
    text = paragraph.text.strip()
    style = paragraph.style.name
    if normalise(text) in {"table of contents", "list of figures", "list of tables"}:
        return 1
    if style == "Heading 1":
        return 1
    if style in {"Heading 2", "Heading 3"}:
        return 2 if not normalise(text).startswith("chapter ") else 1
    if style in {"Heading 4", "Heading 5", "Title"}:
        return 3
    if re.match(r"^\d+\.\d+(?:\.\d+)?\s+[A-Za-z]", text) and len(text) <= 90:
        return 3 if text.count(".") >= 2 else 2
    return None


def cleanup_old_toc(document: Document) -> tuple[object, list[object]]:
    paragraphs = list(document.paragraphs)
    toc_indices = [i for i, p in enumerate(paragraphs) if normalise(p.text) == "table of contents"]
    if not toc_indices:
        raise RuntimeError("Table of Contents heading was not found.")
    first_toc = toc_indices[0]
    figures_index = next(i for i in range(first_toc + 1, len(paragraphs)) if normalise(paragraphs[i].text) == "list of figures")
    anchor = paragraphs[figures_index]
    for paragraph in paragraphs[first_toc:figures_index]:
        paragraph._element.getparent().remove(paragraph._element)
    return anchor, list(document.paragraphs)


def insert_static_clickable_toc(document: Document, anchor) -> None:
    headings = []
    for paragraph in document.paragraphs:
        level = heading_level(paragraph)
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        if level is None or not text or normalise(text) == "table of contents":
            continue
        headings.append((level, text, paragraph))
    bookmark_id = 500
    entries = []
    for position, (level, text, paragraph) in enumerate(headings, 1):
        bookmark_name = f"toc_{position:03d}"
        add_bookmark(paragraph, bookmark_name, bookmark_id)
        bookmark_id += 1
        entries.append((level, text, bookmark_name))

    created = []
    title = document.add_paragraph(style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.5
    add_page_break(title)
    set_tnr(title.add_run("Table of Contents"), 16, bold=True)
    created.append(title._p)
    note = document.add_paragraph(style="Normal")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.line_spacing = 1.5
    set_tnr(note.add_run("Ctrl + Click an entry to open the relevant section."), 10, italic=True)
    created.append(note._p)
    for level, text, bookmark in entries:
        entry = document.add_paragraph(style="Normal")
        make_hyperlink(entry, bookmark, text, {1: 0.0, 2: 0.30, 3: 0.60}.get(level, 0.60))
        created.append(entry._p)
    closing = document.add_paragraph(style="Normal")
    add_page_break(closing)
    created.append(closing._p)
    for element in created:
        anchor._p.addprevious(element)


def append_new_screenshots(document: Document) -> list[tuple[int, str]]:
    screenshots = [
        ("WhatsApp Image 2026-08-04 at 1.07.20 PM.jpeg", "Login Screen"),
        ("WhatsApp Image 2026-08-04 at 1.07.57 PM.jpeg", "Registration Screen"),
        ("WhatsApp Image 2026-08-04 at 1.08.22 PM.jpeg", "SmartWorker Home Page"),
        ("WhatsApp Image 2026-08-04 at 1.09.17 PM.jpeg", "Candidate Dashboard and Job Recommendations"),
        ("WhatsApp Image 2026-08-04 at 1.10.48 PM.jpeg", "Job Categories"),
        ("WhatsApp Image 2026-08-04 at 1.10.48 PM (1).jpeg", "Job Categories - Alternate View"),
        ("WhatsApp Image 2026-08-04 at 1.12.57 PM.jpeg", "Job Details and Application Area"),
        ("WhatsApp Image 2026-08-04 at 1.12.57 PM (1).jpeg", "Job Details - Alternate View"),
        ("WhatsApp Image 2026-08-04 at 1.13.50 PM.jpeg", "Employer Manage Applications"),
        ("WhatsApp Image 2026-08-04 at 1.14.33 PM.jpeg", "Employer Dashboard"),
        ("WhatsApp Image 2026-08-04 at 1.15.41 PM.jpeg", "Post a Job Form"),
        ("WhatsApp Image 2026-08-04 at 1.18.59 PM.jpeg", "Candidate Application History"),
        ("WhatsApp Image 2026-08-04 at 1.21.53 PM.jpeg", "Candidate Profile Update"),
        ("WhatsApp Image 2026-08-04 at 1.24.29 PM.jpeg", "Employer Company Profile Update"),
    ]
    title = document.add_paragraph(style="Heading 1")
    add_page_break(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.5
    set_tnr(title.add_run("APPENDIX A\nSMARTWORKER PROJECT SCREENSHOTS"), 16, bold=True)
    for number, (filename, description) in enumerate(screenshots, 20):
        screen_heading = document.add_paragraph(style="Heading 4")
        add_page_break(screen_heading)
        screen_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        screen_heading.paragraph_format.line_spacing = 1.5
        set_tnr(screen_heading.add_run(description), 13, bold=True)
        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.line_spacing = 1.0
        image_paragraph.add_run().add_picture(str(DOWNLOADS / filename), width=Inches(6.30))
        caption = document.add_paragraph(style="Normal")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.line_spacing = 1.5
        set_tnr(caption.add_run(f"Figure {number}: {description}"), 11)
    return [(number, description) for number, (_filename, description) in enumerate(screenshots, 20)]


def append_figure_list_entries(document: Document, figures: list[tuple[int, str]]) -> None:
    paragraphs = list(document.paragraphs)
    list_tables = next((p for p in paragraphs if normalise(p.text) == "list of tables"), None)
    if list_tables is None:
        return
    for number, description in figures:
        paragraph = document.add_paragraph(style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 1.5
        set_tnr(paragraph.add_run(f"Figure {number}: {description}"), 12)
        list_tables._p.addprevious(paragraph._p)


def fix_known_name_typo(document: Document) -> None:
    for paragraph in document.paragraphs:
        if "Shaheen Shan & Sign:" in paragraph.text:
            for run in paragraph.runs:
                if "Shaheen Shan" in run.text:
                    run.text = run.text.replace("Shaheen Shan", "Shaheen Shah")
                    set_tnr(run)


def validate(docx_path: Path) -> None:
    with zipfile.ZipFile(docx_path) as package:
        if package.testzip() is not None:
            raise RuntimeError("Generated DOCX package is corrupt.")
    document = Document(docx_path)
    if len(document.inline_shapes) < 31:
        raise RuntimeError("Expected existing document images plus 14 appended project screenshots.")
    xml = zipfile.ZipFile(docx_path).read("word/document.xml").decode("utf-8", "ignore")
    if xml.count('w:hyperlink w:anchor="toc_') < 40:
        raise RuntimeError("Clickable Table of Contents links were not created as expected.")


def main() -> None:
    shutil.copy2(SOURCE, FINAL)
    document = Document(FINAL)
    anchor, _paragraphs = cleanup_old_toc(document)
    fix_known_name_typo(document)
    appended_figures = append_new_screenshots(document)
    append_figure_list_entries(document, appended_figures)
    insert_static_clickable_toc(document, anchor)
    format_everything(document)
    document.core_properties.title = "SmartWorker - Final Year Project Thesis"
    document.core_properties.author = "Shaheen Shah and Hamad Khan"
    document.save(FINAL)
    validate(FINAL)
    print(FINAL)


if __name__ == "__main__":
    main()
