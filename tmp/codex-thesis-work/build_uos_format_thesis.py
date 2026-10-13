from __future__ import annotations

import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


WORK = Path(r"D:\My Web Sites\smartworkerapp\tmp\codex-thesis-work")
SOURCE = Path(r"C:\Users\Laptop Valley\Downloads\Thesis format.docx")
FINAL = Path(r"D:\My Web Sites\smartworkerapp\SmartWorker_Thesis_UOS_Format_Shaheen_Shah_Hamad_Khan.docx")
LOGO = WORK / "uos-logo-reference.jpeg"
FIGURES = WORK / "smartworker-figures"
DOWNLOADS = Path(r"C:\Users\Laptop Valley\Downloads")

NAVY = "17365D"
GREY = "666666"
BLACK = "000000"


def w(tag: str) -> str:
    return qn(tag)


def clear_paragraph(paragraph) -> None:
    element = paragraph._element
    for child in list(element):
        if child.tag != w("w:pPr"):
            element.remove(child)


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != w("w:sectPr"):
            body.remove(child)


def remove_template_media(docx_path: Path) -> None:
    """The working copy retains source styles but removes source screenshots before adding SmartWorker media."""
    staged = docx_path.with_suffix(".clean.docx")
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    image_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(staged, "w", zipfile.ZIP_DEFLATED) as output:
        for item in source.infolist():
            if item.filename.startswith("word/media/"):
                continue
            payload = source.read(item.filename)
            if item.filename.startswith("word/_rels/") and item.filename.endswith(".rels"):
                root = ET.fromstring(payload)
                for relation in list(root):
                    if relation.tag == f"{{{rel_ns}}}Relationship" and relation.get("Type") == image_type:
                        root.remove(relation)
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            output.writestr(item, payload)
    shutil.move(staged, docx_path)


def prune_template_parts(docx_path: Path) -> None:
    """Remove obsolete footer/header, hyperlink, and custom-XML parts retained from the source template."""
    staged = docx_path.with_suffix(".pruned.docx")
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    attr_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(docx_path, "r") as source:
        document_root = ET.fromstring(source.read("word/document.xml"))
        used_ids = {
            value
            for node in document_root.iter()
            for attribute, value in node.attrib.items()
            if attribute.startswith(f"{{{attr_ns}}}")
        }
        relationships = ET.fromstring(source.read("word/_rels/document.xml.rels"))
        required = {
            office_rel + "styles",
            office_rel + "numbering",
            office_rel + "settings",
            office_rel + "webSettings",
            office_rel + "theme",
            office_rel + "fontTable",
        }
        retained_targets = set()
        for relation in list(relationships):
            if relation.get("Id") not in used_ids and relation.get("Type") not in required:
                relationships.remove(relation)
                continue
            target = relation.get("Target", "")
            if target and not target.startswith("../") and relation.get("TargetMode") != "External":
                retained_targets.add("word/" + target)
        removed = {
            name
            for name in source.namelist()
            if ((name.startswith("word/footer") or name.startswith("word/header")) and name.endswith(".xml") and name not in retained_targets)
            or name in {"word/footnotes.xml", "word/endnotes.xml"}
            or name.startswith("customXml/")
            or name.startswith("word/_rels/footer")
            or (name.startswith("word/_rels/header") and name.replace("word/_rels/", "word/").replace(".rels", "") not in retained_targets)
        }
        with zipfile.ZipFile(staged, "w", zipfile.ZIP_DEFLATED) as output:
            for item in source.infolist():
                if item.filename in removed:
                    continue
                payload = source.read(item.filename)
                if item.filename == "word/_rels/document.xml.rels":
                    payload = ET.tostring(relationships, encoding="utf-8", xml_declaration=True)
                elif item.filename == "[Content_Types].xml":
                    root = ET.fromstring(payload)
                    for override in list(root):
                        if override.tag == f"{{{content_ns}}}Override" and override.get("PartName", "").lstrip("/") in removed:
                            root.remove(override)
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                elif item.filename == "_rels/.rels":
                    root = ET.fromstring(payload)
                    for relation in list(root):
                        if relation.tag == f"{{{rel_ns}}}Relationship" and relation.get("Target", "").startswith("customXml/"):
                            root.remove(relation)
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                output.writestr(item, payload)
    shutil.move(staged, docx_path)


def set_tnr(run, size: float = 12, bold: bool = False, italic: bool = False, color: str = BLACK) -> None:
    run.font.name = "Times New Roman"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(w(key), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_line_spacing(paragraph, value: float = 1.5) -> None:
    paragraph.paragraph_format.line_spacing = value
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def configure_document(doc: Document) -> None:
    for style_name in ("Normal", "Body Text", "List Paragraph", "Normal (Web)"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        fonts = style._element.rPr.rFonts
        fonts.set(w("w:ascii"), "Times New Roman")
        fonts.set(w("w:hAnsi"), "Times New Roman")
        fonts.set(w("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(0)
    for name, size in (("Heading 1", 16), ("Heading 2", 15), ("Heading 3", 16), ("Heading 4", 14), ("Heading 5", 12)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        fonts = style._element.rPr.rFonts
        fonts.set(w("w:ascii"), "Times New Roman")
        fonts.set(w("w:hAnsi"), "Times New Roman")
        fonts.set(w("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(0.89)
    section.right_margin = Inches(0.69)
    section.top_margin = Inches(1.06)
    section.bottom_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)
    for part in (section.header, section.footer):
        part.is_linked_to_previous = False
        for paragraph in part.paragraphs:
            clear_paragraph(paragraph)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(header, 1.0)
    set_tnr(header.add_run("SMARTWORKER - FINAL YEAR PROJECT THESIS"), 9, bold=True, color=GREY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(footer, 1.0)
    set_tnr(footer.add_run("University of Swat | Department of Software Engineering | Page "), 9, color=GREY)
    add_page_field(footer)

    settings = doc.settings.element
    update = settings.find(w("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(w("w:val"), "true")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(w("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(w("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(w("w:fldCharType"), "separate")
    visible = OxmlElement("w:t")
    visible.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(w("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, visible, end))
    set_tnr(run, 9, color=GREY)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def para(doc: Document, text: str = "", *, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, size: float = 12, bold=False, italic=False, before=0, after=0, first_indent=True):
    paragraph = doc.add_paragraph(style="Body Text")
    paragraph.alignment = alignment
    set_line_spacing(paragraph)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    if first_indent and alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
        paragraph.paragraph_format.first_line_indent = Cm(0.63)
    run = paragraph.add_run(text)
    set_tnr(run, size=size, bold=bold, italic=italic)
    return paragraph


def heading(doc: Document, text: str, level: int = 4, *, alignment=WD_ALIGN_PARAGRAPH.LEFT, page_start=False):
    if page_start:
        page_break(doc)
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = alignment
    set_line_spacing(paragraph)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    size = {1: 16, 2: 15, 3: 16, 4: 14, 5: 12}.get(level, 12)
    set_tnr(run, size=size, bold=True)
    return paragraph


def chapter(doc: Document, number: int, title: str) -> None:
    page_break(doc)
    paragraph = doc.add_paragraph(style="Heading 3")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(paragraph)
    paragraph.paragraph_format.space_before = Pt(35)
    paragraph.paragraph_format.space_after = Pt(18)
    set_tnr(paragraph.add_run(f"CHAPTER {number}\n{title.upper()}"), 16, bold=True)


def caption(doc: Document, number: int, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(paragraph)
    set_tnr(paragraph.add_run(f"Figure {number}: {text}"), 11, bold=False)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(w("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(w("w:fill"), fill)
    node.set(w("w:val"), "clear")


def cell_margin(cell, top=80, start=85, bottom=80, end=85) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    nodes = tc_pr.first_child_found_in("w:tcMar")
    if nodes is None:
        nodes = OxmlElement("w:tcMar")
        tc_pr.append(nodes)
    for side, amount in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = nodes.find(w(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            nodes.append(node)
        node.set(w("w:w"), str(amount))
        node.set(w("w:type"), "dxa")


def table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(w(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(w("w:val"), "single")
        node.set(w("w:sz"), "6")
        node.set(w("w:color"), "000000")
        node.set(w("w:space"), "0")


def table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float]) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    table_borders(tbl)
    all_rows = [tbl.rows[0]]
    for row_data in rows:
        all_rows.append(tbl.add_row())
    for r_index, (row, values) in enumerate(zip(all_rows, [headers] + rows)):
        for cell, value, width_cm in zip(row.cells, values, widths_cm):
            cell.width = Cm(width_cm)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margin(cell)
            if r_index == 0:
                shade(cell, "D9E2F3")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_line_spacing(p)
            p.paragraph_format.first_line_indent = Cm(0)
            set_tnr(p.add_run(value), 10.5, bold=(r_index == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def cover(doc: Document) -> None:
    for _ in range(2):
        para(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(p, 1.0)
    p.add_run().add_picture(str(LOGO), width=Inches(1.15))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(p)
    set_tnr(p.add_run("SMARTWORKER"), 22, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(p)
    set_tnr(p.add_run("AN INTELLIGENT JOB PORTAL AND RECRUITMENT MANAGEMENT SYSTEM"), 14, bold=True)
    para(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(p)
    set_tnr(p.add_run("Submitted by:"), 13, bold=True)
    for person in ("Shaheen Shah (Registration No. ____________________)", "Hamad Khan (Registration No. ____________________)"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_line_spacing(p)
        set_tnr(p.add_run(person), 12)
    para(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(p)
    set_tnr(p.add_run("Supervised by:"), 13, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(p)
    set_tnr(p.add_run("____________________________"), 12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(p)
    set_tnr(p.add_run("Department of Software Engineering"), 13, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(p)
    set_tnr(p.add_run("University of Swat, Khyber Pakhtunkhwa, Pakistan"), 12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(p)
    set_tnr(p.add_run("Session 2022-2026"), 13, bold=True)


def preliminary_pages(doc: Document) -> None:
    heading(doc, "Declaration", 2, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    para(doc, "We, Shaheen Shah and Hamad Khan, declare that the final year project thesis entitled SmartWorker: An Intelligent Job Portal and Recruitment Management System is our own work. The project has been prepared for the award of the Bachelor of Science degree in Software Engineering at the University of Swat. To the best of our knowledge, this work has not been submitted, in whole or in part, to any other institution for the award of a degree.")
    para(doc, "Student Name: Shaheen Shah\t\tSignature: ____________________", alignment=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False, before=18)
    para(doc, "Student Name: Hamad Khan\t\tSignature: ____________________", alignment=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False, before=12)

    heading(doc, "Forwarding Sheet", 2, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    para(doc, "This is to certify that the project entitled SmartWorker: An Intelligent Job Portal and Recruitment Management System has been completed by Shaheen Shah and Hamad Khan under the supervision of the undersigned. The project is submitted in partial fulfilment of the requirements for the Bachelor of Science degree in Software Engineering at the University of Swat.")
    para(doc, "Supervisor: __________________________\nDesignation: _________________________\nDepartment of Software Engineering\nUniversity of Swat", alignment=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False, before=36)

    heading(doc, "Dedication", 2, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    para(doc, "This project is dedicated to our parents, families, teachers, and friends whose encouragement, prayers, and guidance supported us throughout our academic journey.", alignment=WD_ALIGN_PARAGRAPH.CENTER, italic=True, first_indent=False, before=85)

    heading(doc, "Final Approval", 2, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    para(doc, "This is to confirm that the project titled SmartWorker: An Intelligent Job Portal and Recruitment Management System, presented by Shaheen Shah and Hamad Khan, has been reviewed and is accepted towards the completion of the Bachelor of Science degree in Software Engineering at the University of Swat.")
    heading(doc, "Final Project Evaluation Committee", 4, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    for title in ("External Examiner", "Supervisor", "Chairman"):
        para(doc, f"{title}\nName: ____________________________\tDesignation: ____________________________\nInstitute: _________________________\tSignature: _____________________________", alignment=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False, before=10)

    heading(doc, "Acknowledgements", 2, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    para(doc, "First of all, we are thankful to Almighty Allah for giving us the strength, patience, and ability to complete this final year project. We express our sincere gratitude to our teachers and the Department of Software Engineering, University of Swat, for their academic support and guidance. We are also grateful to our parents, friends, and classmates for their encouragement throughout the analysis, design, development, testing, and documentation of SmartWorker.")
    para(doc, "Shaheen Shah and Hamad Khan", alignment=WD_ALIGN_PARAGRAPH.RIGHT, italic=True, first_indent=False, before=18)

    heading(doc, "Abstract", 1, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    para(doc, "SmartWorker is a web-based job portal and recruitment management system designed to connect job candidates, employers, and administrators on one platform. The project addresses the problems of fragmented job discovery, unstructured recruitment records, manual application tracking, and weak communication between applicants and employers. The system provides role-based workflows that allow candidates to manage their profiles and resumes, discover and save jobs, submit applications, receive notifications, and communicate with employers. Employers can create company profiles, post and manage jobs, review applications, identify suitable candidates, and update hiring status. Administrators can moderate users, jobs, contact messages, and applications.")
    para(doc, "The application is implemented using Laravel, PHP, Blade templates, Vite, and Tailwind CSS. SmartWorker also includes an explainable matching service that uses skills, location, and rating information to recommend jobs and candidates. Input validation, role middleware, private document access, password hashing, CSRF protection, and security headers are included to protect the system. The project demonstrates how a centralised web solution can make recruitment more organised, transparent, and accessible.")
    para(doc, "Keywords: Job Portal, Recruitment Management, Laravel, Candidate, Employer, Smart Matching, Web Application.", bold=True, first_indent=False)

    heading(doc, "Table of Contents", 1, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    for entry in ("Declaration", "Forwarding Sheet", "Dedication", "Final Approval", "Acknowledgements", "Abstract", "List of Figures", "Chapter 1: Introduction", "Chapter 2: Software Development Methodology", "Chapter 3: Software Requirement Specification", "Chapter 4: System Design", "Chapter 5: Implementation", "Chapter 6: Testing", "Chapter 7: Conclusion and Future Work", "References", "Appendix A: SmartWorker Screenshots"):
        para(doc, entry, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
    heading(doc, "List of Figures", 1, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    figures = [
        "Figure 1: SmartWorker Use-Case Diagram", "Figure 2: SmartWorker System Architecture", "Figure 3: SmartWorker Entity Relationship Diagram",
        "Figure 4: Login Screen", "Figure 5: Registration Screen", "Figure 6: Home Page", "Figure 7: Candidate Dashboard", "Figure 8: Job Categories", "Figure 9: Job Categories - Alternate View", "Figure 10: Job Details", "Figure 11: Job Details - Alternate View", "Figure 12: Manage Applications", "Figure 13: Employer Dashboard", "Figure 14: Post a Job", "Figure 15: Candidate Application History", "Figure 16: Candidate Profile", "Figure 17: Employer Company Profile",
    ]
    for entry in figures:
        para(doc, entry, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)


def chapter_one(doc: Document) -> None:
    chapter(doc, 1, "Introduction")
    heading(doc, "1.1 Introduction")
    para(doc, "The growth of online recruitment has changed the way candidates search for employment and the way employers find suitable people. However, many recruitment activities are still handled through separate advertisements, social-media posts, spreadsheets, email messages, and manual follow-up. This makes it difficult for candidates to manage their information and for employers to organise job vacancies and applications. SmartWorker was developed as a single web platform to improve these activities.")
    heading(doc, "1.2 Problem Statement")
    para(doc, "Job candidates need a simple and reliable way to create professional profiles, upload resumes, search relevant vacancies, apply for jobs, and check their application status. Employers need an organised system to post jobs, view applications, find candidates, communicate with applicants, and update hiring decisions. A secure role-based system is therefore required to manage these processes without allowing unauthorised access to sensitive information.")
    heading(doc, "1.3 Scope of Project")
    para(doc, "The scope of SmartWorker includes job discovery, candidate and company profiles, job posting, application management, saved jobs, protected CV and resume access, notifications, chat, ratings, categories, contact messages, and administrator moderation. The project supports three major roles: candidate, employer, and administrator. Payroll processing, employee attendance, and a complete enterprise human-resource system are outside the current scope.")
    heading(doc, "1.4 Project Objectives")
    objectives = [
        "To provide a centralised platform for candidates to find and apply for jobs.",
        "To allow employers to post, update, activate, and manage job listings.",
        "To organise candidate profiles, resumes, cover letters, and application status.",
        "To implement role-based access control for candidates, employers, and administrators.",
        "To recommend suitable jobs and candidates using skills, location, and rating data.",
        "To provide notifications, chat, and ratings for improved communication.",
        "To protect user information through validation, authorisation, and private document access.",
    ]
    for index, item in enumerate(objectives, 1):
        para(doc, f"{index}. {item}", alignment=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
    heading(doc, "1.5 Project Features")
    features = [
        ["Candidate", "Profile, skills, resume upload, job search, saved jobs, applications, recommendations, alerts, and chat."],
        ["Employer", "Company profile, job creation, job management, application review, candidate matching, and chat."],
        ["Administrator", "User activation/verification, job moderation, contact-message management, and application oversight."],
        ["Security", "Validation, authentication, role checks, private document downloads, CSRF protection, and security headers."],
    ]
    table(doc, ["Module", "Core Features"], features, [3.8, 13.1])


def chapter_two(doc: Document) -> None:
    chapter(doc, 2, "Software Development Methodology")
    heading(doc, "2.1 Agile Development Model")
    para(doc, "An Agile-oriented software development approach was used for SmartWorker. The project was divided into manageable increments so that requirements, database entities, interfaces, validations, and role workflows could be reviewed throughout development. The iterative approach supports continuous improvement because job portal requirements can be refined while the application is being implemented and tested.")
    heading(doc, "2.2 Project Plan")
    plan = [
        ["Requirements Analysis", "Identify candidate, employer, and administrator workflows."],
        ["Design", "Prepare user-interface structure, database entities, relationships, and access rules."],
        ["Development", "Build authentication, profiles, jobs, applications, dashboards, and administration."],
        ["Smart Features", "Add matching, notifications, chat, ratings, and protected documents."],
        ["Testing", "Run feature tests, validate workflows, and correct defects."],
        ["Documentation", "Prepare project report, diagrams, testing results, and screenshots."],
    ]
    table(doc, ["Phase", "Activities"], plan, [4.4, 12.5])
    heading(doc, "2.3 Selection of Tools and Technologies")
    technologies = [
        ["Programming Language", "PHP 8.3", "Server-side implementation."],
        ["Framework", "Laravel 13", "Routes, controllers, models, migrations, authentication, validation."],
        ["Frontend", "Blade, Vite, Tailwind CSS", "Web pages, layout, assets, and responsive styling."],
        ["Database", "Laravel migration-based relational database", "Users, profiles, jobs, applications, messages, ratings, notifications."],
        ["Testing", "PHPUnit and Laravel Feature Tests", "Authentication, job posting, applications, and route access."],
        ["Development Tools", "VS Code, Composer, npm", "Code editing and dependency management."],
    ]
    table(doc, ["Category", "Technology", "Use"], technologies, [3.5, 4.4, 9.0])
    heading(doc, "2.4 Risk Plan")
    risks = [
        ["Unauthorised role access", "Medium", "Role middleware allows only active and permitted users."],
        ["Invalid input", "Medium", "Form Request validation checks registration, profiles, jobs, and contacts."],
        ["Private file exposure", "High", "Resume/CV download controllers perform authorisation checks."],
        ["Recommendation quality", "Medium", "Explainable scoring can be reviewed and adjusted."],
        ["Deployment configuration", "Medium", "Environment configuration and deployment blueprint are maintained."],
    ]
    table(doc, ["Risk", "Level", "Mitigation"], risks, [4.2, 2.2, 10.5])


def add_center_image(doc: Document, file: Path, width: float) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(paragraph, 1.0)
    paragraph.add_run().add_picture(str(file), width=Inches(width))


def chapter_three(doc: Document) -> None:
    chapter(doc, 3, "Software Requirement Specification")
    heading(doc, "3.1 System Requirements")
    para(doc, "SmartWorker is a web application that requires a computer or mobile device with a modern web browser and an internet connection. For development, the system uses PHP, Laravel, Composer, npm, a relational database, and a code editor. The user-facing interface is designed for candidates, employers, and administrators.")
    heading(doc, "3.2 Functional Requirements")
    functional = [
        ["FR-01", "Registration and login", "Candidates and employers can create accounts, log in, log out, and reset passwords."],
        ["FR-02", "Profile management", "Users can update personal/company details, location, skills, documents, and media."],
        ["FR-03", "Job browsing", "Visitors and candidates can browse active jobs and job categories."],
        ["FR-04", "Job posting", "Employers can create, edit, manage, and delete their job listings."],
        ["FR-05", "Applications", "Candidates can submit a CV and cover letter and view their application history."],
        ["FR-06", "Application handling", "Employers can review applications and update the hiring status."],
        ["FR-07", "Matching and alerts", "The system can recommend jobs/candidates and create job-match notifications."],
        ["FR-08", "Administration", "Administrators can manage users, job listings, contacts, and applications."],
    ]
    table(doc, ["ID", "Function", "Requirement"], functional, [1.5, 4.1, 11.3])
    heading(doc, "3.3 Non-Functional Requirements")
    nonfunctional = [
        ["Usability", "The interface must provide clear navigation and role-specific dashboards."],
        ["Security", "Data must be protected through authentication, authorisation, validation, and secure document access."],
        ["Reliability", "Database constraints and validation must prevent duplicate applications and invalid input."],
        ["Maintainability", "Laravel MVC structure separates routes, controllers, models, services, and requests."],
        ["Scalability", "The system design permits later use of queues, mobile clients, and improved matching services."],
    ]
    table(doc, ["Requirement", "Description"], nonfunctional, [4.0, 12.9])
    heading(doc, "3.4 Use-Case Diagram")
    para(doc, "The use-case diagram shows the main interactions of candidates, employers, and administrators with SmartWorker. Candidate and employer actions are separated through role middleware, while administration is controlled through a dedicated administrator route group.")
    add_center_image(doc, FIGURES / "use-cases.png", 6.2)
    caption(doc, 1, "SmartWorker Use-Case Diagram")


def chapter_four(doc: Document) -> None:
    chapter(doc, 4, "System Design")
    heading(doc, "4.1 System Architecture")
    para(doc, "SmartWorker uses the Model-View-Controller architecture provided by Laravel. Web routes receive requests and send them to controller methods. Controllers manage validation, data access, service calls, and Blade views. Eloquent models communicate with the relational database, while the MatchingService and SmartNotificationService implement important business functions. Private files are stored separately and are downloaded only after access checks.")
    add_center_image(doc, FIGURES / "architecture.png", 6.25)
    caption(doc, 2, "SmartWorker System Architecture")
    heading(doc, "4.2 Entity Relationship Diagram")
    para(doc, "The database design is centred on the User entity. User profiles hold candidate or company details. Employers create job listings, candidates save jobs and create applications, and conversations connect candidates with employers. Ratings and notifications are associated with users and application activity. The diagram below presents the major entities and relationships implemented in the project.")
    add_center_image(doc, FIGURES / "data-model.png", 6.25)
    caption(doc, 3, "SmartWorker Entity Relationship Diagram")
    heading(doc, "4.3 Security and Access Design")
    para(doc, "Access is controlled at route level. Candidate-only actions require a candidate account, employer-only actions require an employer or administrator account, and administrative actions require an administrator account. The middleware also requires the account to be active. Sensitive documents are protected with ownership and role checks before a response is returned.")


def chapter_five(doc: Document) -> None:
    chapter(doc, 5, "Implementation")
    heading(doc, "5.1 Implementation Overview")
    para(doc, "SmartWorker is implemented as a Laravel web application. Named web routes define public pages, authentication, candidate functions, employer functions, shared account controls, and administrator functions. Controllers manage the workflows, Form Requests validate user input, Eloquent models represent database records, and Blade templates present the interface.")
    heading(doc, "5.2 Candidate Module")
    para(doc, "The candidate module provides registration, login, profile update, skills and resume information, job browsing, category browsing, saved jobs, applications, job recommendations, notifications, and chat. The candidate dashboard summarises saved jobs, applications, and recommended vacancies. A candidate can submit one application per job listing and can monitor application status from the application-history screen.")
    heading(doc, "5.3 Employer Module")
    para(doc, "The employer module provides company-profile management, job posting, job management, application review, candidate matching, notifications, and chat. Employers can create a listing by providing job title, location, type, category, salary details, requirements, application information, and company details. The application-management page allows employers to review candidates and update their status to reviewing, shortlisted, hired, or rejected.")
    heading(doc, "5.4 Smart Matching, Notifications, and Chat")
    para(doc, "The matching service calculates a transparent recommendation score from skills, location, and rating information. Candidates receive recommended jobs and employers can view suitable candidates for a listing. The notification service creates job-match alerts for alert-enabled candidates with an appropriate score. The chat module creates conversations between candidates and employers and stores messages within the related conversation.")
    heading(doc, "5.5 Administration and Security")
    para(doc, "The administrator module provides user management, job moderation, contact-message handling, application overview, and account verification. Security controls include password hashing, CSRF protection, Form Request validation, role-based middleware, private resume/CV downloads, and security response headers. The following appendix contains screenshots of the implemented SmartWorker interface.")


def chapter_six(doc: Document) -> None:
    chapter(doc, 6, "Testing")
    heading(doc, "6.1 Testing Approach")
    para(doc, "Testing was carried out through Laravel feature tests and inspection of the implemented workflows. The tests cover authentication, role access, job posting, application workflow, post-job pages, and compatibility routes. Validation rules and database constraints are also used to reject invalid data and prevent duplicate applications or saved-job entries.")
    heading(doc, "6.2 Automated Test Result")
    para(doc, "The project test command was executed during preparation of this thesis. The run completed without a test failure: 13 tests were registered, 5 tests passed, 8 tests were skipped, and 10 assertions were reported. The runner produced a non-blocking warning because it could not update the PHPUnit result-cache file; the warning did not cause a test failure.")
    heading(doc, "6.3 Representative Test Cases")
    tests = [
        ["Authentication", "Register/login with valid credentials", "User is authenticated and directed to the correct role workflow."],
        ["Role Access", "Candidate accesses employer route", "Role middleware rejects unauthorised access."],
        ["Job Posting", "Employer submits valid job data", "Validated job listing is stored and manageable by its owner."],
        ["Application", "Candidate applies for active job", "One application is stored for the candidate/listing pair."],
        ["Documents", "Unauthorised CV/resume request", "Access check prevents public file exposure."],
        ["Notifications", "Eligible job match is created", "Targeted job-match alert can be created for the candidate."],
    ]
    table(doc, ["Area", "Test", "Expected Result"], tests, [3.2, 5.1, 8.6])
    heading(doc, "6.4 Validation")
    para(doc, "The application validates registration information, email uniqueness, password confirmation, job type, coordinates, URLs, salary ranges, required job description, and file types. Resumes are restricted to PDF, DOC, and DOCX formats, while profile and company images are validated before upload. These checks support the reliability and security of the system.")


def chapter_seven(doc: Document) -> None:
    chapter(doc, 7, "Conclusion and Future Work")
    heading(doc, "7.1 Conclusion")
    para(doc, "SmartWorker provides a centralised job portal and recruitment-management system for candidates, employers, and administrators. It combines professional profiles, job listings, applications, saved jobs, protected documents, notifications, chat, ratings, recommendations, and administration in one web application. The role-based architecture, validation, and access-control mechanisms help protect recruitment information and keep user workflows organised.")
    para(doc, "The project meets its objective of simplifying the recruitment process through a practical web system. Its explainable matching approach provides a clear basis for job and candidate recommendations, while the implemented feature tests provide an initial quality baseline for important workflows.")
    heading(doc, "7.2 Future Work")
    future = [
        "Introduce an evaluated machine-learning recommendation model with configurable matching weights.",
        "Develop Android and iOS mobile clients through a secured API.",
        "Add richer hiring analytics, reporting, and application funnel dashboards.",
        "Use queues for high-volume notifications and long-running jobs.",
        "Perform broader browser, accessibility, load, and security testing.",
        "Add audit logs, improved monitoring, and data-retention controls.",
    ]
    for index, item in enumerate(future, 1):
        para(doc, f"{index}. {item}", alignment=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
    heading(doc, "References", 1, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    references = [
        "Laravel. Laravel Documentation. https://laravel.com/docs",
        "PHP Documentation. https://www.php.net/docs.php",
        "Tailwind CSS. Documentation. https://tailwindcss.com/docs",
        "Vite. Guide. https://vite.dev/guide/",
        "PHPUnit. Documentation. https://docs.phpunit.de/",
        "OWASP Foundation. Cheat Sheet Series. https://cheatsheetseries.owasp.org/",
        "SmartWorker project source code: routes, controllers, models, migrations, services, Form Requests, and feature tests.",
    ]
    for reference in references:
        para(doc, reference, first_indent=False)


def screenshot_pages(doc: Document) -> None:
    screenshots = [
        ("WhatsApp Image 2026-08-04 at 1.07.20 PM.jpeg", "Login Screen"),
        ("WhatsApp Image 2026-08-04 at 1.07.57 PM.jpeg", "Registration Screen"),
        ("WhatsApp Image 2026-08-04 at 1.08.22 PM.jpeg", "SmartWorker Home Page"),
        ("WhatsApp Image 2026-08-04 at 1.09.17 PM.jpeg", "Candidate Dashboard and Job Recommendations"),
        ("WhatsApp Image 2026-08-04 at 1.10.48 PM.jpeg", "Job Categories"),
        ("WhatsApp Image 2026-08-04 at 1.10.48 PM (1).jpeg", "Job Categories - Alternate View"),
        ("WhatsApp Image 2026-08-04 at 1.12.57 PM.jpeg", "Job Details and Application Area"),
        ("WhatsApp Image 2026-08-04 at 1.12.57 PM (1).jpeg", "Job Details - Alternate View"),
        ("WhatsApp Image 2026-08-04 at 1.13.50 PM.jpeg", "Employer Manage Applications"),
        ("WhatsApp Image 2026-08-04 at 1.14.33 PM.jpeg", "Employer Dashboard"),
        ("WhatsApp Image 2026-08-04 at 1.15.41 PM.jpeg", "Post a Job Form"),
        ("WhatsApp Image 2026-08-04 at 1.18.59 PM.jpeg", "Candidate Application History"),
        ("WhatsApp Image 2026-08-04 at 1.21.53 PM.jpeg", "Candidate Profile Update"),
        ("WhatsApp Image 2026-08-04 at 1.24.29 PM.jpeg", "Employer Company Profile Update"),
    ]
    heading(doc, "Appendix A", 1, alignment=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    para(doc, "SMARTWORKER IMPLEMENTATION SCREENSHOTS", alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first_indent=False)
    for figure_number, (filename, description) in enumerate(screenshots, 4):
        page_break(doc)
        heading(doc, description, 4, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_center_image(doc, DOWNLOADS / filename, 6.25)
        caption(doc, figure_number, description)


def validate(docx_path: Path) -> None:
    document = Document(docx_path)
    body_text = "\n".join(p.text for p in document.paragraphs)
    forbidden = ("Virtual Study Group", "Sharjeel Muhammad", "Luqman Khan")
    found = [item for item in forbidden if item.lower() in body_text.lower()]
    if found:
        raise RuntimeError(f"Old template content remains: {found}")
    if len(document.inline_shapes) != 18:
        raise RuntimeError(f"Expected 18 images (UOS logo, 3 diagrams, 14 screenshots); found {len(document.inline_shapes)}")
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.text and run.font.name not in (None, "Times New Roman"):
                raise RuntimeError(f"Non-Times New Roman body run detected: {run.font.name}")


def main() -> None:
    if not SOURCE.exists() or not LOGO.exists():
        raise FileNotFoundError("The UOS thesis template or extracted UOS logo is missing.")
    if not all((FIGURES / name).exists() for name in ("architecture.png", "use-cases.png", "data-model.png")):
        raise FileNotFoundError("Required SmartWorker diagrams are missing.")
    shutil.copy2(SOURCE, FINAL)
    document = Document(FINAL)
    clear_body(document)
    document.save(FINAL)
    remove_template_media(FINAL)
    document = Document(FINAL)
    configure_document(document)
    document.core_properties.title = "SmartWorker: An Intelligent Job Portal and Recruitment Management System"
    document.core_properties.author = "Shaheen Shah and Hamad Khan"
    document.core_properties.subject = "Final Year Project Thesis"
    document.core_properties.keywords = "SmartWorker, job portal, recruitment, Laravel"

    cover(document)
    preliminary_pages(document)
    chapter_one(document)
    chapter_two(document)
    chapter_three(document)
    chapter_four(document)
    chapter_five(document)
    chapter_six(document)
    chapter_seven(document)
    screenshot_pages(document)
    document.save(FINAL)
    prune_template_parts(FINAL)
    validate(FINAL)
    with zipfile.ZipFile(FINAL) as package:
        if package.testzip() is not None:
            raise RuntimeError("The generated DOCX package is corrupt.")
    print(FINAL)


if __name__ == "__main__":
    main()
