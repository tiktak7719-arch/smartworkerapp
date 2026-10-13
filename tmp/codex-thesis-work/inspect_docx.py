import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")


def show_paragraphs(document, start, stop):
    for index, paragraph in enumerate(document.paragraphs[start:stop], start):
        text = paragraph.text.replace("\t", " ").strip()
        if text:
            print(f"P{index:04d} | {paragraph.style.name} | {text}")


def show_tables(document):
    for table_index, table in enumerate(document.tables):
        print(f"TABLE {table_index}: {len(table.rows)} rows x {len(table.columns)} cols")
        for row_index, row in enumerate(table.rows):
            values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            print(f"  R{row_index}: {' || '.join(values)}")


def show_summary(document):
    print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} sections={len(document.sections)}")
    for index, section in enumerate(document.sections, 1):
        print(
            f"SECTION {index}: {section.page_width.inches:.2f}x{section.page_height.inches:.2f} "
            f"margins L={section.left_margin.inches:.2f} R={section.right_margin.inches:.2f} "
            f"T={section.top_margin.inches:.2f} B={section.bottom_margin.inches:.2f}"
        )


def main():
    path = Path(sys.argv[1])
    command = sys.argv[2] if len(sys.argv) > 2 else "summary"
    document = Document(path)
    if command == "summary":
        show_summary(document)
    elif command == "paragraphs":
        start = int(sys.argv[3]) if len(sys.argv) > 3 else 0
        stop = int(sys.argv[4]) if len(sys.argv) > 4 else len(document.paragraphs)
        show_paragraphs(document, start, stop)
    elif command == "tables":
        show_tables(document)
    else:
        raise SystemExit(f"Unknown command: {command}")


if __name__ == "__main__":
    main()
