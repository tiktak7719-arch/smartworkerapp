from __future__ import annotations

import re
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


WORK = Path(r"D:\My Web Sites\smartworkerapp\tmp\codex-thesis-work")
sys.path.insert(0, str(WORK))
import build_uos_format_thesis as b


BASE = Path(r"C:\Users\Laptop Valley\Downloads\SmartWorker_Thesis_UOS_Format_Shaheen_Shah_Hamad_Khan.docx")
FINAL = Path(r"D:\My Web Sites\smartworkerapp\SmartWorker_Final_Combined_UOS_Thesis_Shaheen_Shah_Hamad_Khan.docx")
DOWNLOADS = Path(r"C:\Users\Laptop Valley\Downloads")


TOC = [
    ("Declaration", "Declaration", "toc_declaration"),
    ("Forwarding Sheet", "Forwarding Sheet", "toc_forwarding"),
    ("Dedication", "Dedication", "toc_dedication"),
    ("Final Approval", "Final Approval", "toc_approval"),
    ("Acknowledgements", "Acknowledgements", "toc_acknowledgements"),
    ("Abstract", "Abstract", "toc_abstract"),
    ("List of Figures", "List of Figures", "toc_figures"),
    ("CHAPTER 1: INTRODUCTION", "CHAPTER 1 INTRODUCTION", "toc_ch1"),
    ("    1.1 Introduction", "1.1 Introduction", "toc_1_1"),
    ("    1.2 Background and Motivation", "1.2 Background and Motivation", "toc_1_2"),
    ("    1.3 Problem Statement", "1.3 Problem Statement", "toc_1_3"),
    ("    1.4 Scope of Project", "1.4 Scope of Project", "toc_1_4"),
    ("    1.5 Project Objectives", "1.5 Project Objectives", "toc_1_5"),
    ("    1.6 Stakeholders and Benefits", "1.6 Stakeholders and Benefits", "toc_1_6"),
    ("    1.7 Project Deliverables", "1.7 Project Deliverables", "toc_1_7"),
    ("CHAPTER 2: SOFTWARE DEVELOPMENT METHODOLOGY", "CHAPTER 2 SOFTWARE DEVELOPMENT METHODOLOGY", "toc_ch2"),
    ("    2.1 Development Approach", "2.1 Development Approach", "toc_2_1"),
    ("    2.2 Agile Iterations", "2.2 Agile Iterations", "toc_2_2"),
    ("    2.3 Work Breakdown and Schedule", "2.3 Work Breakdown and Schedule", "toc_2_3"),
    ("    2.4 Feasibility Analysis", "2.4 Feasibility Analysis", "toc_2_4"),
    ("    2.5 Tools and Technologies", "2.5 Tools and Technologies", "toc_2_5"),
    ("    2.6 Risk Management", "2.6 Risk Management", "toc_2_6"),
    ("CHAPTER 3: SOFTWARE REQUIREMENT SPECIFICATION", "CHAPTER 3 SOFTWARE REQUIREMENT SPECIFICATION", "toc_ch3"),
    ("    3.1 Stakeholder Requirements", "3.1 Stakeholder Requirements", "toc_3_1"),
    ("    3.2 Functional Requirements", "3.2 Functional Requirements", "toc_3_2"),
    ("    3.3 Non-Functional Requirements", "3.3 Non-Functional Requirements", "toc_3_3"),
    ("    3.4 Use-Case Diagram", "3.4 Use-Case Diagram", "toc_3_4"),
    ("    3.5 Business Rules", "3.5 Business Rules", "toc_3_5"),
    ("    3.6 Inputs and Outputs", "3.6 Inputs and Outputs", "toc_3_6"),
    ("CHAPTER 4: SYSTEM DESIGN", "CHAPTER 4 SYSTEM DESIGN", "toc_ch4"),
    ("    4.1 System Architecture", "4.1 System Architecture", "toc_4_1"),
    ("    4.2 Data Design", "4.2 Data Design", "toc_4_2"),
    ("    4.3 Matching and Notification Design", "4.3 Matching and Notification Design", "toc_4_3"),
    ("    4.4 Security and Access Design", "4.4 Security and Access Design", "toc_4_4"),
    ("    4.5 User Interface Design", "4.5 User Interface Design", "toc_4_5"),
    ("CHAPTER 5: IMPLEMENTATION", "CHAPTER 5 IMPLEMENTATION", "toc_ch5"),
    ("    5.1 Implementation Overview", "5.1 Implementation Overview", "toc_5_1"),
    ("    5.2 Authentication and Profiles", "5.2 Authentication and Profiles", "toc_5_2"),
    ("    5.3 Candidate Module", "5.3 Candidate Module", "toc_5_3"),
    ("    5.4 Employer Module", "5.4 Employer Module", "toc_5_4"),
    ("    5.5 Application Workflow", "5.5 Application Workflow", "toc_5_5"),
    ("    5.6 Communication, Matching, and Alerts", "5.6 Communication, Matching, and Alerts", "toc_5_6"),
    ("    5.7 Administration and Security", "5.7 Administration and Security", "toc_5_7"),
    ("CHAPTER 6: TESTING", "CHAPTER 6 TESTING", "toc_ch6"),
    ("    6.1 Testing Strategy", "6.1 Testing Strategy", "toc_6_1"),
    ("    6.2 Automated Test Results", "6.2 Automated Test Results", "toc_6_2"),
    ("    6.3 Representative Test Cases", "6.3 Representative Test Cases", "toc_6_3"),
    ("    6.4 Validation and User Acceptance", "6.4 Validation and User Acceptance", "toc_6_4"),
    ("CHAPTER 7: CONCLUSION AND FUTURE WORK", "CHAPTER 7 CONCLUSION AND FUTURE WORK", "toc_ch7"),
    ("    7.1 Conclusion", "7.1 Conclusion", "toc_7_1"),
    ("    7.2 Future Work", "7.2 Future Work", "toc_7_2"),
    ("References", "References", "toc_references"),
    ("Appendix A: SmartWorker Implementation Screenshots", "Appendix A", "toc_appendix"),
]


def p(doc, text="", *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, bold=False, italic=False, first=True, before=0):
    return b.para(doc, text, alignment=align, size=size, bold=bold, italic=italic, first_indent=first, before=before)


def h(doc, text, *, level=4, align=WD_ALIGN_PARAGRAPH.LEFT, page_start=False):
    return b.heading(doc, text, level, alignment=align, page_start=page_start)


def cover(doc):
    for _ in range(2):
        p(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, first=False)
    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    b.set_line_spacing(logo, 1.0)
    logo.add_run().add_picture(str(b.LOGO), width=Inches(1.15))
    for text, size, bold in (
        ("SMARTWORKER", 22, True),
        ("AN INTELLIGENT JOB PORTAL AND RECRUITMENT MANAGEMENT SYSTEM", 14, True),
    ):
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        b.set_line_spacing(line)
        b.set_tnr(line.add_run(text), size, bold=bold)
    p(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, first=False)
    for text, size, bold in (
        ("Submitted by:", 13, True),
        ("Shaheen Shah (Registration No. 226022071)", 12, False),
        ("Hamad Khan (Registration No. 226022083)", 12, False),
        ("", 12, False),
        ("Supervised by:", 13, True),
        ("Mr. Abdul Haq", 12, False),
        ("", 12, False),
        ("Department of Software Engineering", 13, True),
        ("University of Swat, Khyber Pakhtunkhwa, Pakistan", 12, False),
        ("Session 2022-2026", 13, True),
    ):
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        b.set_line_spacing(line)
        b.set_tnr(line.add_run(text), size, bold=bold)


def preliminary(doc):
    h(doc, "Declaration", level=2, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    p(doc, "We, Shaheen Shah and Hamad Khan, declare that the final year project thesis entitled SmartWorker: An Intelligent Job Portal and Recruitment Management System is our own work. This thesis is submitted in partial fulfilment of the requirements for the Bachelor of Science degree in Software Engineering at the University of Swat. To the best of our knowledge, this work has not been submitted, in whole or in part, to any other institution for the award of a degree.")
    p(doc, "Shaheen Shah (226022071)\t\tSignature: ____________________", align=WD_ALIGN_PARAGRAPH.LEFT, first=False, before=18)
    p(doc, "Hamad Khan (226022083)\t\tSignature: ____________________", align=WD_ALIGN_PARAGRAPH.LEFT, first=False, before=12)

    h(doc, "Forwarding Sheet", level=2, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    p(doc, "This is to certify that the project entitled SmartWorker: An Intelligent Job Portal and Recruitment Management System has been completed by Shaheen Shah and Hamad Khan under the supervision of Mr. Abdul Haq. The project is submitted in partial fulfilment of the requirements for the Bachelor of Science degree in Software Engineering at the University of Swat.")
    p(doc, "Supervisor: Mr. Abdul Haq\nEmail: abdulhaq@uswat.edu.pk\nDepartment of Software Engineering\nUniversity of Swat", align=WD_ALIGN_PARAGRAPH.LEFT, first=False, before=36)

    h(doc, "Dedication", level=2, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    p(doc, "This project is dedicated to our parents, families, teachers, and friends whose encouragement, prayers, and guidance supported us throughout our academic journey.", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, first=False, before=85)

    h(doc, "Final Approval", level=2, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    p(doc, "This is to confirm that the project titled SmartWorker: An Intelligent Job Portal and Recruitment Management System, presented by Shaheen Shah and Hamad Khan, has been reviewed and is accepted towards the completion of the Bachelor of Science degree in Software Engineering at the University of Swat.")
    h(doc, "Final Project Evaluation Committee", level=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    for title in ("External Examiner", "Supervisor", "Chairman"):
        p(doc, f"{title}\nName: ____________________________\tDesignation: ____________________________\nInstitute: _________________________\tSignature: _____________________________", align=WD_ALIGN_PARAGRAPH.LEFT, first=False, before=10)

    h(doc, "Acknowledgements", level=2, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    p(doc, "First of all, we are thankful to Almighty Allah for giving us the strength, patience, and ability to complete this final year project. We express our sincere gratitude to our supervisor, Mr. Abdul Haq, and to the Department of Software Engineering, University of Swat, for academic guidance and support. We are also grateful to our parents, friends, and classmates for their encouragement during the analysis, design, development, testing, and documentation of SmartWorker.")
    p(doc, "Shaheen Shah and Hamad Khan", align=WD_ALIGN_PARAGRAPH.RIGHT, italic=True, first=False, before=18)

    h(doc, "Abstract", level=1, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    p(doc, "SmartWorker is a web-based job portal and recruitment management system designed to connect candidates, employers, and administrators through one organised platform. The project addresses common recruitment problems including scattered job advertisements, weak application tracking, unstructured applicant records, delayed communication, and limited visibility of relevant opportunities. SmartWorker enables candidates to create professional profiles, upload resumes, browse and save active jobs, apply with a cover letter, track their application status, receive notifications, and communicate with employers. Employers can maintain company profiles, post vacancies, review candidate applications, manage hiring status, and identify suitable candidates.")
    p(doc, "The application is implemented as a Laravel web application using PHP, Blade templates, Vite, and Tailwind CSS. Its architecture follows a structured Model-View-Controller approach and includes routes, controllers, models, migrations, validation requests, services, and feature tests. An explainable matching service recommends jobs and candidates using skills, location, and rating information. In the candidate recommendation score, skills contribute 60%, location 30%, and rating 10%; employer-side candidate recommendations use skills 65%, location 25%, and rating 10%. This approach keeps recommendations understandable for users and can be improved in later versions.")
    p(doc, "The system includes authentication, role-based access control, account-status checks, validation, password hashing, CSRF protection, protected resume/CV download routes, notifications, chat, ratings, categories, and administrator moderation. SmartWorker demonstrates how a secure, centralised, and user-friendly platform can improve the recruitment workflow for both job seekers and employers.")
    p(doc, "Keywords: Job Portal, Recruitment Management, Laravel, Candidate, Employer, Smart Matching, Web Application.", bold=True, first=False)

    h(doc, "Table of Contents", level=1, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    note = p(doc, "Hold Ctrl and click an entry to go directly to that section.", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, first=False)
    note.paragraph_format.space_after = Pt(4)
    for display, _target, _anchor in TOC:
        p(doc, display, align=WD_ALIGN_PARAGRAPH.LEFT, first=False)

    h(doc, "List of Figures", level=1, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    figures = [
        "Figure 1: SmartWorker Use-Case Diagram", "Figure 2: SmartWorker System Architecture",
        "Figure 3: SmartWorker Entity Relationship Diagram", "Figure 4: SmartWorker Matching and Alert Workflow",
        "Figure 5: Login Screen", "Figure 6: Registration Screen", "Figure 7: Home Page",
        "Figure 8: Candidate Dashboard", "Figure 9: Job Categories", "Figure 10: Job Categories - Alternate View",
        "Figure 11: Job Details and Application Area", "Figure 12: Job Details - Alternate View",
        "Figure 13: Employer Manage Applications", "Figure 14: Employer Dashboard", "Figure 15: Post a Job Form",
        "Figure 16: Candidate Application History", "Figure 17: Candidate Profile Update", "Figure 18: Employer Company Profile Update",
    ]
    for item in figures:
        p(doc, item, align=WD_ALIGN_PARAGRAPH.LEFT, first=False)


def chapter_one(doc):
    b.chapter(doc, 1, "Introduction")
    h(doc, "1.1 Introduction")
    p(doc, "Online recruitment has changed the way people search for employment and how organisations advertise vacancies. In many local and small-scale recruitment settings, however, job information is still dispersed across social-media posts, printed notices, separate web pages, spreadsheets, email messages, and manual telephone follow-up. Candidates may miss suitable vacancies or lose track of applications, while employers may find it difficult to organise job posts, applications, candidate documents, and hiring decisions. SmartWorker was developed as a unified platform to make these activities more visible, structured, and manageable.")
    h(doc, "1.2 Background and Motivation")
    p(doc, "The need for a single recruitment workspace is especially important for students, fresh graduates, skilled workers, and small employers who need a practical way to present opportunities and professional information. A useful portal should do more than show job advertisements: it should support profile creation, search, job categories, secure applications, application status, communication, and responsible handling of documents. The project is motivated by the idea that candidates and employers should have clear, role-specific workflows without needing multiple unrelated tools.")
    h(doc, "1.3 Problem Statement")
    p(doc, "Job candidates need a reliable way to create professional profiles, upload resumes, search relevant vacancies, apply for jobs, and check their application progress. Employers need an organised method to create vacancies, review applicants, record hiring decisions, and communicate with appropriate candidates. Without a central system, recruitment becomes time-consuming, less transparent, and more likely to contain duplicated records or missed communication. A secure role-based web application is therefore required to manage these processes while protecting users and their documents.")
    h(doc, "1.4 Scope of Project")
    p(doc, "The scope of SmartWorker includes job discovery, candidate and company profiles, job posting, application management, saved jobs, resume/CV access, notifications, chat, ratings, categories, contact messages, and administrator moderation. The platform supports candidate, employer, and administrator roles. Payroll processing, employee attendance, tax management, and a complete enterprise human-resource information system are outside the current scope.")
    h(doc, "1.5 Project Objectives")
    objectives = [
        "Provide a centralised platform through which candidates can find and apply for jobs.",
        "Allow employers to create, update, activate, and manage job listings and company information.",
        "Organise candidate profiles, skills, resumes, cover letters, application records, and status updates.",
        "Implement controlled access for candidates, employers, and administrators.",
        "Recommend suitable jobs and candidates using transparent skills, location, and rating information.",
        "Provide notifications, chat, and ratings to improve communication and engagement.",
        "Protect recruitment information through validation, authorisation, and private document access.",
    ]
    for i, item in enumerate(objectives, 1):
        p(doc, f"{i}. {item}", align=WD_ALIGN_PARAGRAPH.LEFT, first=False)
    h(doc, "1.6 Stakeholders and Benefits")
    b.table(doc, ["Stakeholder", "Benefit from SmartWorker"], [
        ["Candidates", "Professional profile, vacancy browsing, saved jobs, applications, alerts, and clear status tracking."],
        ["Employers", "Company profile, structured job posting, application review, matching support, and recruitment communication."],
        ["Administrators", "Moderation of users, jobs, contacts, and applications through dedicated management screens."],
        ["University / Community", "A practical example of secure, role-based web engineering and organised recruitment workflows."],
    ], [4.3, 12.6])
    h(doc, "1.7 Project Deliverables")
    p(doc, "The principal deliverables are the SmartWorker Laravel web application, its database migrations and models, role-specific dashboards, responsive Blade views, matching and notification services, automated feature tests, project diagrams, this thesis, and implementation screenshots. The screenshots in Appendix A provide evidence of the candidate and employer functionality implemented in the system.")


def chapter_two(doc):
    b.chapter(doc, 2, "Software Development Methodology")
    h(doc, "2.1 Development Approach")
    p(doc, "SmartWorker was developed using an Agile-oriented approach. The work was organised into small, testable increments so requirements and user flows could be revisited while the project was progressing. This approach is suitable for a job portal because interface needs, application rules, and matching behaviour become clearer as candidate, employer, and administrator workflows are implemented.")
    h(doc, "2.2 Agile Iterations")
    b.table(doc, ["Iteration", "Main Outcome"], [
        ["1. Analysis", "Identification of users, recruitment problems, role permissions, and core requirements."],
        ["2. Foundation", "Laravel project structure, database migrations, authentication, and profile foundations."],
        ["3. Recruitment Core", "Job browsing, job posting, applications, saved jobs, and dashboards."],
        ["4. Smart Features", "Matching, notifications, chat, ratings, and protected document handling."],
        ["5. Quality", "Validation, access checks, feature tests, screenshot evidence, and documentation."],
    ], [4.0, 12.9])
    h(doc, "2.3 Work Breakdown and Schedule")
    p(doc, "The project was planned through the stages of requirement analysis, design, implementation, testing, and documentation. Requirement analysis focused on role needs and application rules. Design included data relationships, routes, page flows, and interface structure. Implementation built the Laravel modules in manageable groups, followed by testing and documentation. This staged plan made it possible to keep functional features connected to their corresponding design and test evidence.")
    b.table(doc, ["Work Package", "Activities and Evidence"], [
        ["Requirements", "Problem definition, scope, user roles, functional and non-functional requirements."],
        ["Design", "Use cases, database relationships, MVC architecture, matching workflow, and access rules."],
        ["Development", "Profiles, job listings, applications, statuses, dashboards, notifications, chat, and administration."],
        ["Testing", "Feature tests, input validation, workflow inspection, and result recording."],
        ["Documentation", "Thesis chapters, diagrams, figures, references, and implementation screenshots."],
    ], [4.0, 12.9])
    h(doc, "2.4 Feasibility Analysis")
    b.table(doc, ["Area", "Assessment"], [
        ["Technical", "Laravel, PHP, Blade, Vite, and Tailwind CSS provide an appropriate web stack for role-based recruitment workflows."],
        ["Operational", "Candidates, employers, and administrators use focused pages and dashboards that match their normal tasks."],
        ["Economic", "Open-source framework and development tools reduce the cost of prototyping and academic development."],
        ["Schedule", "Iterative development supports delivery of core recruitment features before optional smart enhancements."],
    ], [4.0, 12.9])
    h(doc, "2.5 Tools and Technologies")
    b.table(doc, ["Category", "Technology", "Use in Project"], [
        ["Programming", "PHP 8.3", "Server-side implementation."],
        ["Framework", "Laravel 13", "Routes, controllers, models, migrations, validation, authentication, and tests."],
        ["Frontend", "Blade, Vite, Tailwind CSS", "Responsive server-rendered pages, layouts, and assets."],
        ["Database", "Laravel migration-based relational database", "Users, profiles, jobs, applications, messages, ratings, and notifications."],
        ["Testing", "PHPUnit and Laravel feature tests", "Authentication, route access, job posting, and application workflow tests."],
        ["Development", "VS Code, Composer, npm", "Code editing, dependency management, and front-end assets."],
    ], [3.2, 4.4, 9.3])
    h(doc, "2.6 Risk Management")
    b.table(doc, ["Risk", "Level", "Mitigation"], [
        ["Unauthorised role access", "Medium", "Role middleware permits only active users with required roles."],
        ["Invalid or incomplete input", "Medium", "Form Request validation checks profile, job, contact, and application data."],
        ["Private file exposure", "High", "Resume and CV download routes perform ownership and role checks."],
        ["Poor recommendation relevance", "Medium", "Explainable weights can be reviewed and tuned after feedback."],
        ["Deployment misconfiguration", "Medium", "Environment configuration and deployment settings are maintained with the project."],
    ], [4.2, 2.2, 10.5])


def chapter_three(doc):
    b.chapter(doc, 3, "Software Requirement Specification")
    h(doc, "3.1 Stakeholder Requirements")
    p(doc, "Candidates require a simple way to maintain their profile, identify relevant jobs, submit applications, save useful vacancies, receive alerts, and follow their application progress. Employers require structured publishing of vacancies, an overview of applications, candidate review, status updates, and company information. Administrators require tools to supervise site activity and manage users, listings, contact messages, and applications. These needs were translated into role-specific routes, controllers, views, and data relationships.")
    h(doc, "3.2 Functional Requirements")
    b.table(doc, ["ID", "Function", "Requirement"], [
        ["FR-01", "Accounts", "Candidates and employers can register, log in, log out, reset passwords, and use their role-specific accounts."],
        ["FR-02", "Profiles", "Users can update personal or company details, location, skills, documents, and profile media."],
        ["FR-03", "Job Discovery", "Visitors and candidates can browse active jobs, categories, and job details."],
        ["FR-04", "Job Posting", "Employers can create, edit, manage, and remove job listings that belong to them."],
        ["FR-05", "Applications", "Candidates can submit a CV and cover letter and view application history and status."],
        ["FR-06", "Application Review", "Employers can review applications and set review, shortlist, hire, or reject status."],
        ["FR-07", "Matching and Alerts", "The system recommends jobs/candidates and creates job-match notifications for eligible candidates."],
        ["FR-08", "Administration", "Administrators can manage users, listings, contact records, and applications."],
    ], [1.5, 3.8, 11.6])
    h(doc, "3.3 Non-Functional Requirements")
    b.table(doc, ["Requirement", "Description"], [
        ["Usability", "Clear navigation, labelled forms, and separate dashboards should support role-specific tasks."],
        ["Security", "Authentication, authorisation, validation, private document access, and secure request handling must protect data."],
        ["Reliability", "Validation and database constraints must prevent invalid input and duplicate applications or saved-job entries."],
        ["Maintainability", "Laravel MVC separates routes, controllers, models, services, views, and validation requests."],
        ["Scalability", "The design permits later queues, APIs, mobile clients, analytics, and improved recommendation models."],
    ], [4.0, 12.9])
    h(doc, "3.4 Use-Case Diagram")
    p(doc, "The use-case diagram shows the main interactions of candidates, employers, and administrators with SmartWorker. Candidate and employer actions are separated through role middleware, while administrative actions are restricted to a dedicated administrator route group.")
    b.add_center_image(doc, b.FIGURES / "use-cases.png", 6.2)
    b.caption(doc, 1, "SmartWorker Use-Case Diagram")
    h(doc, "3.5 Business Rules")
    b.table(doc, ["Rule", "Explanation"], [
        ["Role separation", "Candidate, employer, and administrator routes require the applicable role and an active account."],
        ["Job ownership", "An employer may manage only job listings associated with that employer."],
        ["Application uniqueness", "A candidate may submit only one application for the same job listing."],
        ["Private documents", "Resume/CV downloads require an authorised user and cannot be used as public links."],
        ["Active opportunities", "Job search and matching use active job listings."],
    ], [4.0, 12.9])
    h(doc, "3.6 Inputs and Outputs")
    p(doc, "Important inputs include registration details, profile information, skills, location, job title, category, salary information, job description, application CV, cover letter, and application decisions. Important outputs include job listings, personalised recommendations, application history, employer application lists, notification alerts, chat messages, and administration records. Validation is applied before accepted input is written to the database.")


def chapter_four(doc):
    b.chapter(doc, 4, "System Design")
    h(doc, "4.1 System Architecture")
    p(doc, "SmartWorker uses the Model-View-Controller architecture provided by Laravel. Web routes receive requests and send them to controller methods. Controllers coordinate validation, data access, service calls, and Blade views. Eloquent models communicate with the relational database, while the MatchingService and SmartNotificationService handle recommendation and alert functions. Sensitive files are stored separately and returned only after access checks.")
    b.add_center_image(doc, b.FIGURES / "architecture.png", 6.25)
    b.caption(doc, 2, "SmartWorker System Architecture")
    h(doc, "4.2 Data Design")
    p(doc, "The data model is centred on the User entity. User profiles contain candidate or company data. Employers create job listings, candidates save jobs and create applications, and conversations connect candidates with employers. Ratings and notifications are associated with user activity. This structure keeps recruitment records linked to the correct account and supports the role-specific pages shown in the implementation screenshots.")
    b.add_center_image(doc, b.FIGURES / "data-model.png", 6.25)
    b.caption(doc, 3, "SmartWorker Entity Relationship Diagram")
    h(doc, "4.3 Matching and Notification Design")
    p(doc, "SmartWorker uses an explainable scoring approach instead of an opaque recommendation model. When jobs are recommended to a candidate, skills contribute 60 percent, location contributes 30 percent, and rating contributes 10 percent. When candidates are recommended to an employer, skills contribute 65 percent, location contributes 25 percent, and rating contributes 10 percent. The score can be understood by users and provides a practical base for future machine-learning improvements. Eligible candidates can receive job-match notifications through the notification service.")
    b.add_center_image(doc, b.FIGURES / "matching-workflow.png", 6.0)
    b.caption(doc, 4, "SmartWorker Matching and Alert Workflow")
    h(doc, "4.4 Security and Access Design")
    p(doc, "Access is controlled at route level. Candidate-only actions require a candidate account, employer functions require an employer or administrator account where appropriate, and administration requires an administrator account. The middleware also checks that the account is active. Form Requests validate values before use, passwords are hashed, CSRF protection is provided by Laravel, and documents are served through authorised controller actions. Security response headers provide an additional browser-side protection layer.")
    h(doc, "4.5 User Interface Design")
    p(doc, "The interface uses a common navigation structure with role-based actions. Public users can search jobs and access account pages. Candidate dashboards focus on saved jobs, applications, and recommendations. Employer dashboards focus on active jobs, received applications, and job posting. Forms group related details into clear sections, such as basic information, job details, company information, and application information. The design is implemented through Blade templates and Tailwind CSS assets.")


def chapter_five(doc):
    b.chapter(doc, 5, "Implementation")
    h(doc, "5.1 Implementation Overview")
    p(doc, "SmartWorker is implemented as a Laravel web application. Named web routes define public pages, authentication, candidate functions, employer functions, shared account controls, and administrator functions. Controllers manage workflows, Form Requests validate user input, Eloquent models represent database records, services implement matching and notifications, and Blade templates render the interface. This separation improves maintainability and keeps business rules close to their responsible components.")
    h(doc, "5.2 Authentication and Profiles")
    p(doc, "Users can create an account as a candidate or employer, log in, log out, and use password-reset functionality. After authentication, each user sees role-appropriate navigation and dashboard actions. Candidates maintain personal profile information, skills, location, resume, and professional details. Employers maintain company information and related organisation details. The registration, login, candidate profile, and employer profile screens are shown in Appendix A.")
    h(doc, "5.3 Candidate Module")
    p(doc, "The candidate module supports job browsing, category browsing, saved jobs, applications, application history, profile updates, recommended jobs, notifications, and chat. A candidate can open an active job, review its description and details, attach an application CV, add a cover letter, and later check the recorded application status. The dashboard brings together key information so a candidate can manage their job-seeking process from one account.")
    h(doc, "5.4 Employer Module")
    p(doc, "The employer module supports company profile management, job posting, job management, application review, candidate matching, notifications, and chat. Job forms capture job title, location, type, category, salary, description, application details, and company information. The employer can review received applications and update their status to reviewing, shortlisted, hired, or rejected. These activities appear in the employer dashboard, post-job form, and manage-applications screens in Appendix A.")
    h(doc, "5.5 Application Workflow")
    p(doc, "The application workflow begins when a candidate selects an active vacancy and submits a CV and cover letter. The system records the association between the candidate and the job listing. The candidate can then see the application in application history. The employer reviews received applications, can add an optional note, and updates the status. This workflow replaces informal application tracking with a consistent record available to the correct users.")
    h(doc, "5.6 Communication, Matching, and Alerts")
    p(doc, "SmartWorker includes chat conversations between candidates and employers, notifications for relevant job matches, and ratings. The matching service uses available skills, location, and rating information to produce recommendations, while the notification service can create an alert for candidates who have enabled relevant job alerts. These supporting features make the portal more than a static job-listing website by encouraging timely communication and discovery of relevant opportunities.")
    h(doc, "5.7 Administration and Security")
    p(doc, "The administrator module provides user management, job moderation, contact-message handling, application overview, and account verification. The security implementation includes password hashing, CSRF protection, Form Request validation, role middleware, active-account requirements, private resume/CV downloads, and response security headers. Together these controls ensure that each type of user performs only the actions permitted for their role.")


def chapter_six(doc):
    b.chapter(doc, 6, "Testing")
    h(doc, "6.1 Testing Strategy")
    p(doc, "Testing combined automated Laravel feature tests with inspection of the implemented recruitment workflows. The focus was placed on authentication, access permissions, job posting, application processing, validation, and compatibility routes. The test strategy also considered the manual interface flow: candidates should be able to discover and apply for opportunities, and employers should be able to create jobs and process received applications.")
    h(doc, "6.2 Automated Test Results")
    p(doc, "The project test command was executed during preparation of this thesis. The run completed without a test failure: 13 tests were registered, 5 tests passed, 8 tests were skipped, and 10 assertions were reported. The runner produced a non-blocking warning because it could not update the PHPUnit result-cache file; the warning did not cause a test failure. Skipped tests remain a useful target for future test-suite expansion.")
    h(doc, "6.3 Representative Test Cases")
    b.table(doc, ["Area", "Test", "Expected Result"], [
        ["Authentication", "Register or log in with valid credentials", "The user is authenticated and directed to a valid role workflow."],
        ["Role Access", "Candidate attempts employer-only route", "Role middleware rejects unauthorised access."],
        ["Job Posting", "Employer submits valid job data", "A validated job listing is stored and can be managed by its owner."],
        ["Application", "Candidate applies for an active job", "One application is stored for the candidate and job listing pair."],
        ["Documents", "Unauthorised CV/resume request", "The access check prevents public file exposure."],
        ["Notifications", "Eligible job match is created", "A targeted job-match alert can be created for the candidate."],
    ], [3.1, 5.2, 8.6])
    h(doc, "6.4 Validation and User Acceptance")
    p(doc, "The application validates registration information, email uniqueness, password confirmation, job type, coordinates, URLs, salary ranges, required job descriptions, and file types. Resumes are restricted to PDF, DOC, and DOCX formats, while profile and company images are validated before upload. Manual acceptance was supported by reviewing the implemented screens for registration, login, job search, categories, dashboards, application management, job posting, and profile updates. The Appendix A screenshots provide visible evidence of these workflows.")


def chapter_seven(doc):
    b.chapter(doc, 7, "Conclusion and Future Work")
    h(doc, "7.1 Conclusion")
    p(doc, "SmartWorker provides a centralised job portal and recruitment management system for candidates, employers, and administrators. It brings together professional profiles, job listings, applications, saved jobs, protected documents, notifications, chat, ratings, recommendations, and administration in one Laravel web application. The role-based architecture, input validation, and access-control mechanisms help protect recruitment information and keep user workflows organised.")
    p(doc, "The project achieves its main objective of simplifying recruitment through a practical, accessible web solution. Its explainable matching approach gives a clear starting point for job and candidate recommendations, while the implemented feature tests provide an initial quality baseline for core workflows. The screenshots included in this thesis demonstrate the visible implementation of the principal modules.")
    h(doc, "7.2 Future Work")
    for i, item in enumerate([
        "Introduce an evaluated machine-learning recommendation model with configurable matching weights.",
        "Develop Android and iOS mobile clients through a secured API.",
        "Add hiring analytics, reporting, and application-funnel dashboards.",
        "Use queues for high-volume notifications and other long-running processes.",
        "Perform broader browser, accessibility, load, and security testing.",
        "Add audit logs, monitoring, and clearer data-retention controls.",
    ], 1):
        p(doc, f"{i}. {item}", align=WD_ALIGN_PARAGRAPH.LEFT, first=False)
    h(doc, "References", level=1, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    for reference in [
        "Laravel. Laravel Documentation. https://laravel.com/docs",
        "PHP Documentation. https://www.php.net/docs.php",
        "Tailwind CSS. Documentation. https://tailwindcss.com/docs",
        "Vite. Guide. https://vite.dev/guide/",
        "PHPUnit. Documentation. https://docs.phpunit.de/",
        "OWASP Foundation. Cheat Sheet Series. https://cheatsheetseries.owasp.org/",
        "SmartWorker project source code: routes, controllers, models, migrations, services, Form Requests, and feature tests.",
    ]:
        p(doc, reference, first=False)


def appendix(doc):
    screenshots = [
        ("WhatsApp Image 2026-08-04 at 1.07.20 PM.jpeg", "Login Screen"),
        ("WhatsApp Image 2026-08-04 at 1.07.57 PM.jpeg", "Registration Screen"),
        ("WhatsApp Image 2026-08-04 at 1.08.22 PM.jpeg", "SmartWorker Home Page"),
        ("WhatsApp Image 2026-08-04 at 1.09.17 PM.jpeg", "Candidate Dashboard and Job Recommendations"),
        ("WhatsApp Image 2026-08-04 at 1.10.48 PM.jpeg", "Job Categories"),
        ("WhatsApp Image 2026-08-04 at 1.10.48 PM (1).jpeg", "Job Categories - Alternate View"),
        ("WhatsApp Image 2026-08-04 at 1.12.57 PM.jpeg", "Job Details and Application Area"),
        ("WhatsApp Image 2026-08-04 at 1.12.57 PM (1).jpeg", "Job Details - Alternate View"),
        ("WhatsApp Image 2026-08-04 at 1.13.50 PM.jpeg", "Employer Manage Applications"),
        ("WhatsApp Image 2026-08-04 at 1.14.33 PM.jpeg", "Employer Dashboard"),
        ("WhatsApp Image 2026-08-04 at 1.15.41 PM.jpeg", "Post a Job Form"),
        ("WhatsApp Image 2026-08-04 at 1.18.59 PM.jpeg", "Candidate Application History"),
        ("WhatsApp Image 2026-08-04 at 1.21.53 PM.jpeg", "Candidate Profile Update"),
        ("WhatsApp Image 2026-08-04 at 1.24.29 PM.jpeg", "Employer Company Profile Update"),
    ]
    h(doc, "Appendix A", level=1, align=WD_ALIGN_PARAGRAPH.CENTER, page_start=True)
    p(doc, "SMARTWORKER IMPLEMENTATION SCREENSHOTS", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first=False)
    for number, (filename, description) in enumerate(screenshots, 5):
        b.page_break(doc)
        h(doc, description, align=WD_ALIGN_PARAGRAPH.CENTER)
        b.add_center_image(doc, DOWNLOADS / filename, 6.25)
        b.caption(doc, number, description)


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def add_bookmark(paragraph, name: str, identifier: int):
    element = paragraph._p
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    element.insert(1 if element.pPr is not None else 0, start)
    element.append(end)


def internal_link(paragraph, text: str, anchor: str):
    element = paragraph._p
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), "Times New Roman")
    props.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    props.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    element.append(hyperlink)


def make_clickable_toc(path: Path):
    document = Document(path)
    paragraphs = document.paragraphs
    toc_heading = next(i for i, p_ in enumerate(paragraphs) if normalize(p_.text) == "table of contents")
    note_index = toc_heading + 1
    start = note_index + 1
    toc_indexes = set(range(start, start + len(TOC)))
    for index, (_display, target, anchor) in enumerate(TOC, 100):
        candidates = [
            paragraph for paragraph_index, paragraph in enumerate(paragraphs)
            if paragraph_index not in toc_indexes and normalize(paragraph.text) == normalize(target)
        ]
        if not candidates:
            raise RuntimeError(f"TOC target not found: {target}")
        add_bookmark(candidates[-1], anchor, index)
    entries = paragraphs[start:start + len(TOC)]
    if len(entries) != len(TOC):
        raise RuntimeError("TOC entry count does not match.")
    for paragraph, (display, _target, anchor) in zip(entries, TOC):
        internal_link(paragraph, display, anchor)
    document.save(path)


def validate(path: Path):
    document = Document(path)
    if len(document.inline_shapes) != 19:
        raise RuntimeError(f"Expected 19 images; found {len(document.inline_shapes)}")
    required = ["Shaheen Shah (Registration No. 226022071)", "Hamad Khan (Registration No. 226022083)", "Mr. Abdul Haq", "Laravel 13"]
    contents = "\n".join(
        [p_.text for p_ in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    missing = [item for item in required if item not in contents]
    if missing:
        raise RuntimeError(f"Required combined data missing: {missing}")
    with zipfile.ZipFile(path) as package:
        if package.testzip() is not None:
            raise RuntimeError("DOCX package is corrupt.")


def main():
    if not BASE.exists():
        raise FileNotFoundError(f"Base UOS-format thesis missing: {BASE}")
    b.SOURCE = BASE
    b.FINAL = FINAL
    b.cover = cover
    b.preliminary_pages = preliminary
    b.chapter_one = chapter_one
    b.chapter_two = chapter_two
    b.chapter_three = chapter_three
    b.chapter_four = chapter_four
    b.chapter_five = chapter_five
    b.chapter_six = chapter_six
    b.chapter_seven = chapter_seven
    b.screenshot_pages = appendix
    b.validate = validate
    b.main()
    make_clickable_toc(FINAL)
    validate(FINAL)
    print(FINAL)


if __name__ == "__main__":
    main()
